#!/usr/bin/env python3
"""
Calcula sentimiento semanal para Facebook, Twitter, YouTube y medios usando
diccionarios de polaridad y la ponderacion de redes planteada en el script
original.

Salida:
    - Un solo archivo de Excel
    - Una sola hoja
    - Una fila por semana ISO
    - Columnas con:
        1. sentimiento ponderado de redes
        2. sentimiento de medios
        3. promedio de ambos
"""

from __future__ import annotations

import argparse
import re
import unicodedata
from datetime import date, timedelta
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


# =============================================================================
# CONFIGURACION
# =============================================================================
BASE_DIR = Path(__file__).resolve().parents[2]
TEXTOS_DIR = BASE_DIR / "data" / "text" / "radar_weekly_flat"
DICT_DIR = BASE_DIR / "data" / "reference" / "dictionaries_nlp" / "diccionarios_polaridad"
OUTPUT_PATH = (
    BASE_DIR
    / "data"
    / "processed"
    / "modeling"
    / "aceptacion_digital_redes_medios_sentimiento_semanal.xlsx"
)

FUENTES = {
    "facebook": TEXTOS_DIR / "facebook_semana_texto",
    "twitter": TEXTOS_DIR / "twitter_semana_texto",
    "youtube": TEXTOS_DIR / "youtube_semana_texto",
    "medios": TEXTOS_DIR / "medios_semana_texto",
}
ALL_SOURCES = tuple(FUENTES.keys())

STOPLIST_PATH = DICT_DIR / "stop_list_espanol_limpia.txt"
DICT_POS_PATH = DICT_DIR / "diccionario_palabras_positivas.txt"
DICT_NEG_PATH = DICT_DIR / "diccionario_palabras_negativas.txt"

# Parametros del modelo de ponderacion del script original
U_FACEBOOK = 93.0
U_TWITTER = 16.9
U_YOUTUBE = 84.0

F_FACEBOOK = 0.40
F_TWITTER = 1.25
F_YOUTUBE = 1.00  # Escenario moderado del script original

ARCHIVO_RE = re.compile(r"^(?P<anio>\d{2})_(?P<semana>\d{2})_(?P<fuente>[a-z]+)\.txt$")
ARCHIVO_CANONICO_RE = re.compile(r"^(?P<start>\d{4}-\d{2}-\d{2})_(?P<fuente>[a-z]+)\.txt$")
PERIODO_ISO_RE = re.compile(r"^(?P<anio>\d{4})-W(?P<semana>\d{2})$")
TOKEN_RE = re.compile(r"[a-z]+")

HOJA_EXCEL = "sentimiento_semanal"
COLUMNAS_SENTIMIENTO = [
    "sentimiento_facebook",
    "sentimiento_twitter",
    "sentimiento_youtube",
    "sentimiento_redes_ponderado",
    "sentimiento_medios",
    "promedio_redes_medios",
]
COLUMNAS_PRINCIPALES = [
    "anio_iso",
    "semana_iso",
    "periodo_iso",
    "sentimiento_facebook",
    "sentimiento_twitter",
    "sentimiento_youtube",
    "sentimiento_redes_ponderado",
    "sentimiento_medios",
    "promedio_redes_medios",
]


# =============================================================================
# FUNCIONES BASE
# =============================================================================
def normalizar(texto: str) -> str:
    texto = texto.lower().strip().replace("\ufeff", "")
    nfkd = unicodedata.normalize("NFKD", texto)
    return "".join(char for char in nfkd if not unicodedata.combining(char))


def cargar_lista(path: Path) -> set[str]:
    palabras = set()
    with path.open(encoding="utf-8") as handle:
        for linea in handle:
            palabra = normalizar(linea.replace("\t", " ").strip())
            if palabra:
                palabras.add(palabra)
    return palabras


def tokenizar_y_limpiar(texto: str, stopwords: set[str]) -> list[str]:
    tokens = TOKEN_RE.findall(normalizar(texto))
    return [token for token in tokens if len(token) > 2 and token not in stopwords]


def calcular_pesos_redes() -> dict[str, float]:
    n_facebook = U_FACEBOOK * F_FACEBOOK
    n_twitter = U_TWITTER * F_TWITTER
    n_youtube = U_YOUTUBE * F_YOUTUBE
    total = n_facebook + n_twitter + n_youtube

    return {
        "facebook": n_facebook / total,
        "twitter": n_twitter / total,
        "youtube": n_youtube / total,
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Calcula sentimiento semanal canónico. "
            "Sin flags recompone todo; con filtros actualiza incrementalmente semanas/fuentes objetivo."
        )
    )
    parser.add_argument(
        "--sources",
        nargs="+",
        choices=ALL_SOURCES,
        default=list(ALL_SOURCES),
        help="Fuentes a recalcular. Default: todas.",
    )
    parser.add_argument(
        "--week",
        dest="weeks",
        action="append",
        default=[],
        help="Periodo ISO objetivo YYYY-Www. Repite el flag para varios periodos.",
    )
    parser.add_argument("--date-from", help="Fecha inicial para seleccionar semanas objetivo. Formato YYYY-MM-DD.")
    parser.add_argument("--date-to", help="Fecha final para seleccionar semanas objetivo. Formato YYYY-MM-DD.")
    parser.add_argument(
        "--full-refresh",
        action="store_true",
        help="Fuerza recomposición completa e ignora merge incremental.",
    )
    return parser.parse_args()


def parse_periodo_iso(value: str) -> tuple[int, int]:
    match = PERIODO_ISO_RE.match(value.strip())
    if not match:
        raise ValueError(f"Periodo ISO inválido: {value}. Usa YYYY-Www.")
    return int(match.group("anio")), int(match.group("semana"))


def construir_periodos_objetivo(
    weeks: list[str],
    date_from_value: str | None,
    date_to_value: str | None,
) -> set[tuple[int, int]]:
    periodos = {parse_periodo_iso(value) for value in weeks}

    if date_from_value or date_to_value:
        if not date_from_value or not date_to_value:
            raise ValueError("Debes proporcionar --date-from y --date-to juntos.")
        try:
            start_date = date.fromisoformat(date_from_value)
            end_date = date.fromisoformat(date_to_value)
        except ValueError as exc:
            raise ValueError("Las fechas deben usar el formato YYYY-MM-DD.") from exc
        if end_date < start_date:
            raise ValueError("--date-to no puede ser anterior a --date-from.")

        cursor = start_date
        while cursor <= end_date:
            anio_iso, semana_iso, _ = cursor.isocalendar()
            periodos.add((anio_iso, semana_iso))
            cursor += timedelta(days=1)

    return periodos


def extraer_periodo_iso(path: Path, fuente_esperada: str) -> tuple[tuple[int, int], int]:
    match = ARCHIVO_RE.match(path.name)
    if match:
        fuente = match.group("fuente")
        if fuente != fuente_esperada:
            raise ValueError(
                f"Se esperaba fuente '{fuente_esperada}' pero el archivo es '{path.name}'"
            )

        anio_iso = 2000 + int(match.group("anio"))
        semana_iso = int(match.group("semana"))
        return (anio_iso, semana_iso), 1

    match = ARCHIVO_CANONICO_RE.match(path.name)
    if not match:
        raise ValueError(f"Nombre de archivo no valido: {path.name}")

    fuente = match.group("fuente")
    if fuente != fuente_esperada:
        raise ValueError(
            f"Se esperaba fuente '{fuente_esperada}' pero el archivo es '{path.name}'"
        )

    inicio_semana = date.fromisoformat(match.group("start"))
    anio_iso, semana_iso, _ = inicio_semana.isocalendar()
    return (anio_iso, semana_iso), 2


def indexar_archivos(directorio: Path, fuente: str) -> dict[tuple[int, int], Path]:
    if not directorio.exists():
        raise FileNotFoundError(f"No existe el directorio: {directorio}")

    archivos: dict[tuple[int, int], tuple[int, Path]] = {}
    for path in sorted(directorio.glob("*.txt")):
        periodo, prioridad = extraer_periodo_iso(path, fuente)
        previo = archivos.get(periodo)
        if previo is None:
            archivos[periodo] = (prioridad, path)
            continue
        if prioridad > previo[0]:
            archivos[periodo] = (prioridad, path)
            continue
        if prioridad == previo[0]:
            raise ValueError(
                f"Periodo duplicado para {fuente}: {periodo} en {path} y {previo[1]}"
            )
    return {periodo: payload[1] for periodo, payload in archivos.items()}


def analizar_archivo(
    path: Path,
    stopwords: set[str],
    dict_pos: set[str],
    dict_neg: set[str],
) -> dict[str, float | int]:
    texto = path.read_text(encoding="utf-8", errors="ignore")
    tokens = tokenizar_y_limpiar(texto, stopwords)

    positivas = sum(token in dict_pos for token in tokens)
    negativas = sum(token in dict_neg for token in tokens)
    total_sentimiento = positivas + negativas

    if total_sentimiento:
        sentimiento = (positivas - negativas) / total_sentimiento
    else:
        sentimiento = 0.0

    return {
        "tokens_limpios": len(tokens),
        "positivas": int(positivas),
        "negativas": int(negativas),
        "total_sentimiento": int(total_sentimiento),
        "sentimiento": float(sentimiento),
    }


def promedio_ponderado(valores: dict[str, float], pesos: dict[str, float]) -> float:
    disponibles = {
        fuente: valor
        for fuente, valor in valores.items()
        if valor is not None and not pd.isna(valor)
    }
    if not disponibles:
        return float("nan")

    peso_total = sum(pesos[fuente] for fuente in disponibles)
    if peso_total == 0:
        return float("nan")

    return sum(disponibles[fuente] * pesos[fuente] for fuente in disponibles) / peso_total


def promedio_simple(*valores: float) -> float:
    disponibles = [valor for valor in valores if valor is not None and not pd.isna(valor)]
    if not disponibles:
        return float("nan")
    return sum(disponibles) / len(disponibles)


def construir_fila(
    periodo: tuple[int, int],
    analisis_por_fuente: dict[str, dict[str, float | int] | None],
    pesos_redes: dict[str, float],
) -> dict[str, float | int | str]:
    anio_iso, semana_iso = periodo

    fila: dict[str, float | int | str] = {
        "anio_iso": anio_iso,
        "semana_iso": semana_iso,
        "periodo_iso": f"{anio_iso}-W{semana_iso:02d}",
    }

    for fuente in FUENTES:
        analisis = analisis_por_fuente.get(fuente)
        if analisis is None:
            fila[f"sentimiento_{fuente}"] = float("nan")
            continue

        fila[f"sentimiento_{fuente}"] = analisis["sentimiento"]

    sentimiento_redes = promedio_ponderado(
        {
            "facebook": fila["sentimiento_facebook"],
            "twitter": fila["sentimiento_twitter"],
            "youtube": fila["sentimiento_youtube"],
        },
        pesos_redes,
    )

    fila["sentimiento_redes_ponderado"] = sentimiento_redes
    fila["promedio_redes_medios"] = promedio_simple(
        sentimiento_redes,
        fila["sentimiento_medios"],
    )

    return fila


def normalizar_dataframe(dataframe: pd.DataFrame) -> pd.DataFrame:
    if dataframe.empty:
        return pd.DataFrame(columns=COLUMNAS_PRINCIPALES)
    normalizado = dataframe.copy()
    for columna in COLUMNAS_PRINCIPALES:
        if columna not in normalizado.columns:
            normalizado[columna] = pd.NA
    return normalizado[COLUMNAS_PRINCIPALES]


def recalcular_metricas_fila(fila: pd.Series, pesos_redes: dict[str, float]) -> pd.Series:
    sentimiento_redes = promedio_ponderado(
        {
            "facebook": fila.get("sentimiento_facebook"),
            "twitter": fila.get("sentimiento_twitter"),
            "youtube": fila.get("sentimiento_youtube"),
        },
        pesos_redes,
    )
    fila["sentimiento_redes_ponderado"] = sentimiento_redes
    fila["promedio_redes_medios"] = promedio_simple(
        sentimiento_redes,
        fila.get("sentimiento_medios"),
    )
    return fila


def fusionar_actualizacion_parcial(
    existing_df: pd.DataFrame,
    partial_df: pd.DataFrame,
    fuentes_actualizadas: set[str],
    pesos_redes: dict[str, float],
) -> pd.DataFrame:
    base = normalizar_dataframe(existing_df).set_index("periodo_iso", drop=False)
    partial = normalizar_dataframe(partial_df).set_index("periodo_iso", drop=False)

    for periodo_iso, row in partial.iterrows():
        if periodo_iso not in base.index:
            base.loc[periodo_iso, COLUMNAS_PRINCIPALES] = row.reindex(COLUMNAS_PRINCIPALES)
        else:
            base.loc[periodo_iso, "anio_iso"] = row["anio_iso"]
            base.loc[periodo_iso, "semana_iso"] = row["semana_iso"]
            base.loc[periodo_iso, "periodo_iso"] = row["periodo_iso"]
            for fuente in fuentes_actualizadas:
                base.loc[periodo_iso, f"sentimiento_{fuente}"] = row[f"sentimiento_{fuente}"]
        base.loc[periodo_iso] = recalcular_metricas_fila(base.loc[periodo_iso], pesos_redes)

    return (
        normalizar_dataframe(base.reset_index(drop=True))
        .sort_values(["anio_iso", "semana_iso"])
        .reset_index(drop=True)
    )


def ajustar_hoja_excel(writer: pd.ExcelWriter, dataframe: pd.DataFrame) -> None:
    hoja = writer.sheets[HOJA_EXCEL]
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions

    encabezado_fill = PatternFill("solid", fgColor="1F4E78")
    encabezado_font = Font(color="FFFFFF", bold=True)

    for cell in hoja[1]:
        cell.fill = encabezado_fill
        cell.font = encabezado_font

    columnas_sentimiento = {col: idx + 1 for idx, col in enumerate(dataframe.columns)}

    for column_cells in hoja.columns:
        letra = get_column_letter(column_cells[0].column)
        max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        hoja.column_dimensions[letra].width = min(max(max_len + 2, 12), 28)

    for nombre_columna in COLUMNAS_SENTIMIENTO:
        indice = columnas_sentimiento.get(nombre_columna)
        if not indice:
            continue
        letra = get_column_letter(indice)
        for row in range(2, hoja.max_row + 1):
            hoja[f"{letra}{row}"].number_format = "0.0000"


def guardar_excel(dataframe: pd.DataFrame, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        dataframe.to_excel(writer, sheet_name=HOJA_EXCEL, index=False)
        ajustar_hoja_excel(writer, dataframe)


# =============================================================================
# EJECUCION
# =============================================================================
def main() -> None:
    args = parse_args()
    print("=" * 72)
    print("SENTIMIENTO SEMANAL DE REDES Y MEDIOS")
    print("=" * 72)

    stopwords = cargar_lista(STOPLIST_PATH)
    dict_pos = cargar_lista(DICT_POS_PATH)
    dict_neg = cargar_lista(DICT_NEG_PATH)
    pesos_redes = calcular_pesos_redes()

    print("\nDiccionarios cargados:")
    print(f"  Stopwords: {len(stopwords):,}")
    print(f"  Positivas: {len(dict_pos):,}")
    print(f"  Negativas: {len(dict_neg):,}")

    print("\nPesos de redes:")
    print(f"  Facebook: {pesos_redes['facebook']:.4f}")
    print(f"  Twitter:  {pesos_redes['twitter']:.4f}")
    print(f"  YouTube:  {pesos_redes['youtube']:.4f}")
    print(f"  f_youtube usado: {F_YOUTUBE:.2f}")

    fuentes_seleccionadas = tuple(args.sources)
    periodos_objetivo = construir_periodos_objetivo(args.weeks, args.date_from, args.date_to)

    archivos_por_fuente = {
        fuente: indexar_archivos(directorio, fuente)
        for fuente, directorio in FUENTES.items()
    }

    for fuente, archivos in archivos_por_fuente.items():
        print(f"  {fuente.title():<8}: {len(archivos):>3} archivos")

    if periodos_objetivo:
        periodos = sorted(periodos_objetivo)
    else:
        periodos = sorted(
            {
                periodo
                for archivos in archivos_por_fuente.values()
                for periodo in archivos
            }
        )

    filas = []
    for anio_iso, semana_iso in periodos:
        periodo = (anio_iso, semana_iso)
        analisis_por_fuente = {}
        faltantes = []

        for fuente, archivos in archivos_por_fuente.items():
            if fuente not in fuentes_seleccionadas:
                analisis_por_fuente[fuente] = None
                continue
            path = archivos.get(periodo)
            if path is None:
                analisis_por_fuente[fuente] = None
                faltantes.append(fuente)
                continue
            analisis_por_fuente[fuente] = analizar_archivo(path, stopwords, dict_pos, dict_neg)

        if faltantes:
            print(
                f"Aviso: faltan fuentes en {anio_iso}-W{semana_iso:02d}: "
                + ", ".join(faltantes)
            )

        filas.append(construir_fila(periodo, analisis_por_fuente, pesos_redes))

    dataframe = pd.DataFrame(filas).sort_values(["anio_iso", "semana_iso"]).reset_index(drop=True)
    dataframe = normalizar_dataframe(dataframe)

    full_refresh = args.full_refresh or (not periodos_objetivo and set(fuentes_seleccionadas) == set(ALL_SOURCES))
    if full_refresh:
        print("\nModo de ejecución: recomposición completa")
        final_df = dataframe
    else:
        print("\nModo de ejecución: actualización incremental")
        if OUTPUT_PATH.exists():
            existing_df = pd.read_excel(OUTPUT_PATH, sheet_name=HOJA_EXCEL)
        else:
            existing_df = pd.DataFrame(columns=COLUMNAS_PRINCIPALES)
        final_df = fusionar_actualizacion_parcial(
            existing_df,
            dataframe,
            set(fuentes_seleccionadas),
            pesos_redes,
        )
        print(f"  Fuentes actualizadas: {', '.join(fuentes_seleccionadas)}")
        if periodos_objetivo:
            print(f"  Semanas objetivo: {len(periodos_objetivo):,}")

    guardar_excel(final_df, OUTPUT_PATH)

    print("\nResumen:")
    print(f"  Semanas calculadas en esta ejecución: {len(dataframe):,}")
    print(f"  Semanas totales en salida: {len(final_df):,}")
    print(f"  Archivo generado: {OUTPUT_PATH}")
    print("=" * 72)


if __name__ == "__main__":
    main()
