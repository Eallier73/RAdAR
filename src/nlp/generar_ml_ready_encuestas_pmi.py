#!/usr/bin/env python3
"""
Refresca el dataset semanal ML-ready a partir de:

1. un scaffold semanal de encuestas/proyecciones ya estabilizado
2. el consolidado PMI normalizado V5/V10

El resultado vuelve a materializar:
    data/processed/modeling/ml_ready_monica_villarreal_encuestas_pmi_1.xlsx

Este script resuelve el tramo operativo:
    PMI normalizado -> ML-ready semanal con features PMI

No reemplaza la reconstruccion posterior:
    datos_ml_0.xlsx -> dataset maestro final
"""

from __future__ import annotations

import argparse
import re
import unicodedata
from datetime import date, timedelta
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    from .week_resolution import parse_week_token, normalize_to_iso_week_start
except ImportError:  # pragma: no cover - fallback para ejecucion directa
    from week_resolution import parse_week_token, normalize_to_iso_week_start


BASE_DIR = Path(__file__).resolve().parents[2]
DEFAULT_SCAFFOLD_PATH = (
    BASE_DIR / "data" / "processed" / "modeling" / "ml_ready_monica_villarreal_encuestas_pmi_1.xlsx"
)
DEFAULT_PMI_PATH = (
    BASE_DIR
    / "data"
    / "reference"
    / "dictionaries_nlp"
    / "resultados_clasificacion_temas"
    / "pmi_confianza_corpus_unido_consolidado_normalizado.xlsx"
)
DEFAULT_OUTPUT_PATH = DEFAULT_SCAFFOLD_PATH

TRAIN_SHEET = "ML_Ready_Train"
ALL_WEEKS_SHEET = "ML_Ready_AllWeeks"
README_SHEET = "README"
PMI_SHEETS = {
    "Consolidado_V5": "v5",
    "Consolidado_V10": "v10",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Genera/refresca el ML-ready semanal uniendo scaffold de encuestas con PMI normalizado.",
    )
    parser.add_argument(
        "--scaffold",
        type=Path,
        default=DEFAULT_SCAFFOLD_PATH,
        help=f"Workbook base con la grilla semanal de encuestas/proyecciones. Default: {DEFAULT_SCAFFOLD_PATH}",
    )
    parser.add_argument(
        "--pmi-normalized",
        type=Path,
        default=DEFAULT_PMI_PATH,
        help=f"Consolidado PMI normalizado. Default: {DEFAULT_PMI_PATH}",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT_PATH,
        help=f"Archivo de salida ML-ready. Default: {DEFAULT_OUTPUT_PATH}",
    )
    parser.add_argument(
        "--through-date",
        type=str,
        default=None,
        help="Fecha maxima inclusiva YYYY-MM-DD. Si se usa, el PMI se reconstruye solo hasta esa semana.",
    )
    return parser.parse_args()


def parse_optional_date(value: str | None) -> date | None:
    if not value:
        return None
    try:
        return date.fromisoformat(value)
    except ValueError as exc:
        raise SystemExit(f"Fecha invalida para --through-date: {value!r}. Usa YYYY-MM-DD.") from exc


def normalize_header(value: str) -> str:
    text = str(value).strip().lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = text.replace("\n", "_").replace(" ", "_")
    text = re.sub(r"_+", "_", text).strip("_")
    return text


def normalize_date_column(series: pd.Series) -> pd.Series:
    if pd.api.types.is_datetime64_any_dtype(series):
        return pd.to_datetime(series)
    if pd.api.types.is_numeric_dtype(series):
        base = pd.Timestamp("1899-12-30")
        return base + pd.to_timedelta(pd.to_numeric(series), unit="D")
    return pd.to_datetime(series, errors="coerce")


def ensure_week_metadata(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out["fecha_inicio_semana"] = normalize_date_column(out["fecha_inicio_semana"])
    out["fecha_inicio_semana"] = out["fecha_inicio_semana"].dt.normalize()

    if "iso_year" not in out.columns or "iso_week" not in out.columns:
        iso = out["fecha_inicio_semana"].dt.isocalendar()
        out["iso_year"] = iso["year"].astype(int)
        out["iso_week"] = iso["week"].astype(int)
    else:
        out["iso_year"] = pd.to_numeric(out["iso_year"], errors="raise").astype(int)
        out["iso_week"] = pd.to_numeric(out["iso_week"], errors="raise").astype(int)

    out["semana_iso"] = out["iso_year"].astype(str) + "-W" + out["iso_week"].map(lambda value: f"{value:02d}")
    out["iso_yearweek_num"] = out["iso_year"] * 100 + out["iso_week"]
    return out


def load_pmi_sheet(path: Path, sheet_name: str, prefix: str, *, through_date: date | None) -> pd.DataFrame:
    raw = pd.read_excel(path, sheet_name=sheet_name)
    raw = raw[raw.iloc[:, 0].astype(str).str.upper() != "TOTAL"].copy()

    week_starts = [parse_week_token(value) for value in raw.iloc[:, 0]]
    out = pd.DataFrame(
        {
            "fecha_inicio_semana": pd.to_datetime(week_starts),
        }
    )
    iso = out["fecha_inicio_semana"].dt.isocalendar()
    out["iso_year"] = iso["year"].astype(int)
    out["iso_week"] = iso["week"].astype(int)
    out["semana_iso"] = out["iso_year"].astype(str) + "-W" + out["iso_week"].map(lambda value: f"{value:02d}")
    out["iso_yearweek_num"] = out["iso_year"] * 100 + out["iso_week"]

    for column in raw.columns[1:]:
        target_name = f"{prefix}_{normalize_header(column)}"
        out[target_name] = pd.to_numeric(raw[column], errors="coerce")

    if through_date:
        cutoff = pd.Timestamp(normalize_to_iso_week_start(through_date))
        out = out[out["fecha_inicio_semana"] <= cutoff].copy()

    return out.sort_values(["iso_year", "iso_week"]).reset_index(drop=True)


def load_pmi_features(path: Path, *, through_date: date | None) -> pd.DataFrame:
    parts = [
        load_pmi_sheet(path, sheet_name, prefix, through_date=through_date)
        for sheet_name, prefix in PMI_SHEETS.items()
    ]

    merged = parts[0]
    for part in parts[1:]:
        merged = merged.merge(
            part,
            on=["fecha_inicio_semana", "iso_year", "iso_week", "semana_iso", "iso_yearweek_num"],
            how="outer",
            validate="one_to_one",
        )
    return merged.sort_values(["iso_year", "iso_week"]).reset_index(drop=True)


def extend_scaffold(allweeks: pd.DataFrame, *, through_date: date | None) -> tuple[pd.DataFrame, int]:
    if through_date is None:
        return allweeks, 0

    out = ensure_week_metadata(allweeks)
    target_week_start = pd.Timestamp(normalize_to_iso_week_start(through_date))
    if out.empty or out["fecha_inicio_semana"].max() >= target_week_start:
        return out, 0

    appended_rows: list[pd.Series] = []
    current = out.sort_values(["iso_year", "iso_week"]).reset_index(drop=True)
    next_week_start = current.iloc[-1]["fecha_inicio_semana"] + timedelta(days=7)

    while next_week_start <= target_week_start:
        row = current.iloc[-1].copy()
        iso_year, iso_week, _ = next_week_start.date().isocalendar()
        row["fecha_inicio_semana"] = next_week_start
        row["semana_iso"] = f"{iso_year}-W{iso_week:02d}"
        row["iso_year"] = iso_year
        row["iso_week"] = iso_week
        row["iso_yearweek_num"] = iso_year * 100 + iso_week
        appended_rows.append(row)
        current = pd.concat([current, pd.DataFrame([row])], ignore_index=True)
        next_week_start += timedelta(days=7)

    return current, len(appended_rows)


def ordered_pmi_columns(scaffold: pd.DataFrame, pmi_df: pd.DataFrame) -> list[str]:
    scaffold_cols = [col for col in scaffold.columns if col.startswith("v5_") or col.startswith("v10_")]
    pmi_cols = [col for col in pmi_df.columns if col.startswith("v5_") or col.startswith("v10_")]
    if scaffold_cols:
        missing = [col for col in scaffold_cols if col not in pmi_cols]
        if missing:
            raise SystemExit(f"Faltan columnas PMI esperadas en el consolidado normalizado: {missing}")
        return scaffold_cols
    return sorted(pmi_cols)


def merge_scaffold_with_pmi(scaffold: pd.DataFrame, pmi_df: pd.DataFrame) -> pd.DataFrame:
    scaffold = ensure_week_metadata(scaffold)
    pmi_cols = ordered_pmi_columns(scaffold, pmi_df)

    base = scaffold.drop(columns=pmi_cols + ["has_full_pmi_features"], errors="ignore").copy()
    merged = base.merge(
        pmi_df[["iso_year", "iso_week", *pmi_cols]],
        on=["iso_year", "iso_week"],
        how="left",
        validate="one_to_one",
    )
    merged["has_full_pmi_features"] = merged[pmi_cols].notna().all(axis=1).astype(int)

    original_cols = list(scaffold.columns)
    has_full_idx = original_cols.index("has_full_pmi_features") if "has_full_pmi_features" in original_cols else len(base.columns)
    prefix = [col for col in original_cols[:has_full_idx] if col not in pmi_cols and col != "has_full_pmi_features"]
    suffix = [col for col in original_cols[has_full_idx + 1 :] if col not in pmi_cols and col != "has_full_pmi_features"]
    ordered = prefix + ["has_full_pmi_features"] + pmi_cols + suffix

    return merged[ordered].sort_values(["iso_year", "iso_week"]).reset_index(drop=True)


def build_train_sheet(allweeks: pd.DataFrame) -> pd.DataFrame:
    mask = allweeks["has_full_pmi_features"].eq(1)
    for column in ("target_serie_3e", "target_serie_4e"):
        if column in allweeks.columns:
            mask &= allweeks[column].notna()
    return allweeks.loc[mask].copy().reset_index(drop=True)


def build_readme(
    *,
    scaffold_path: Path,
    pmi_path: Path,
    through_date: date | None,
    rows_train: int,
    rows_allweeks: int,
    appended_weeks: int,
    weeks_full_pmi: int,
) -> pd.DataFrame:
    rows = [
        ("Objetivo", "Scaffold semanal de encuestas/proyecciones enriquecido con PMI normalizado V5/V10."),
        ("Scaffold fuente", str(scaffold_path)),
        ("PMI normalizado", str(pmi_path)),
        ("Hoja principal", TRAIN_SHEET),
        ("Hoja secundaria", ALL_WEEKS_SHEET),
        ("Clave de union", "iso_year + iso_week"),
        ("Through date", through_date.isoformat() if through_date else "sin limite"),
        ("Filas ML_Ready_Train", rows_train),
        ("Filas ML_Ready_AllWeeks", rows_allweeks),
        ("Semanas con PMI completo", weeks_full_pmi),
        ("Semanas extendidas desde scaffold", appended_weeks),
        (
            "Nota metodologica",
            "Si el through-date rebasa el scaffold base, se extiende semanalmente por carry-forward de la ultima fila disponible antes de injertar el PMI.",
        ),
    ]
    return pd.DataFrame(rows, columns=["Campo", "Valor"])


def apply_format(writer: pd.ExcelWriter, sheet_name: str, dataframe: pd.DataFrame) -> None:
    sheet = writer.sheets[sheet_name]
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = font

    for column_cells in sheet.columns:
        letter = get_column_letter(column_cells[0].column)
        max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        sheet.column_dimensions[letter].width = min(max(max_len + 2, 12), 32)

    number_formats = {
        "fecha_inicio_semana": "yyyy-mm-dd",
        "mitofsky": "0.0",
        "mitofsky_locf": "0.0",
        "scr": "0.0",
        "scr_locf": "0.0",
        "demoscopia": "0.0",
        "demoscopia_locf": "0.0",
        "rubrum_eval_general": "0.00",
        "rubrum_eval_general_locf": "0.00",
        "rubrum_x10": "0.0",
        "rubrum_x10_locf": "0.0",
        "consenso_mensual_3e": "0.000",
        "consenso_mensual_3e_locf": "0.000",
        "consenso_mensual_4e": "0.000",
        "consenso_mensual_4e_locf": "0.000",
        "target_serie_3e": "0.000",
        "target_serie_4e": "0.000",
    }

    for column in dataframe.columns:
        if column.startswith("v5_") or column.startswith("v10_"):
            number_formats[column] = "0.0000"

    columns = {column: idx + 1 for idx, column in enumerate(dataframe.columns)}
    for column_name, number_format in number_formats.items():
        index = columns.get(column_name)
        if not index:
            continue
        letter = get_column_letter(index)
        for row in range(2, sheet.max_row + 1):
            sheet[f"{letter}{row}"].number_format = number_format


def write_output(path: Path, train_df: pd.DataFrame, allweeks_df: pd.DataFrame, readme_df: pd.DataFrame) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        train_df.to_excel(writer, sheet_name=TRAIN_SHEET, index=False)
        allweeks_df.to_excel(writer, sheet_name=ALL_WEEKS_SHEET, index=False)
        readme_df.to_excel(writer, sheet_name=README_SHEET, index=False)
        apply_format(writer, TRAIN_SHEET, train_df)
        apply_format(writer, ALL_WEEKS_SHEET, allweeks_df)
        apply_format(writer, README_SHEET, readme_df)


def main() -> None:
    args = parse_args()
    through_date = parse_optional_date(args.through_date)

    print("=" * 72)
    print("GENERACION DE ML_READY DESDE PMI NORMALIZADO")
    print("=" * 72)
    print(f"Scaffold:    {args.scaffold}")
    print(f"PMI input:   {args.pmi_normalized}")
    print(f"Output:      {args.output}")
    print(f"Through:     {through_date.isoformat() if through_date else '(sin limite)'}")

    scaffold_allweeks = pd.read_excel(args.scaffold, sheet_name=ALL_WEEKS_SHEET)
    scaffold_allweeks, appended_weeks = extend_scaffold(scaffold_allweeks, through_date=through_date)
    pmi_features = load_pmi_features(args.pmi_normalized, through_date=through_date)
    merged_allweeks = merge_scaffold_with_pmi(scaffold_allweeks, pmi_features)
    train_df = build_train_sheet(merged_allweeks)

    readme_df = build_readme(
        scaffold_path=args.scaffold,
        pmi_path=args.pmi_normalized,
        through_date=through_date,
        rows_train=len(train_df),
        rows_allweeks=len(merged_allweeks),
        appended_weeks=appended_weeks,
        weeks_full_pmi=int(merged_allweeks["has_full_pmi_features"].sum()),
    )
    write_output(args.output, train_df, merged_allweeks, readme_df)

    print(f"Filas {TRAIN_SHEET}: {len(train_df):,}")
    print(f"Filas {ALL_WEEKS_SHEET}: {len(merged_allweeks):,}")
    print(f"Semanas con PMI completo: {int(merged_allweeks['has_full_pmi_features'].sum()):,}")
    print(f"Semanas extendidas: {appended_weeks}")
    print(f"Archivo generado: {args.output}")
    print("=" * 72)


if __name__ == "__main__":
    main()
